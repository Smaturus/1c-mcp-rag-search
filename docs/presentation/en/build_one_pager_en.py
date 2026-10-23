"""Build the one-page English A4 summary (PPTX source, exported to PDF).

Run:  python docs/presentation/en/build_one_pager_en.py
Writes: docs/presentation/en/one-pager-en.pptx
Then convert it to docs/presentation/en/ONE_PAGER_EN.pdf with LibreOffice Kit.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "one-pager-en.pptx"

# A4 portrait, inches
W = Inches(8.27)
H = Inches(11.69)
M = Inches(0.5)

NAVY = RGBColor(0x0F, 0x3D, 0x5C)
ACCENT = RGBColor(0x1B, 0x7F, 0xAE)
LIGHT = RGBColor(0xEA, 0xF2, 0xF7)
INK = RGBColor(0x1F, 0x2A, 0x33)
GRAY = RGBColor(0x5B, 0x67, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RULE = RGBColor(0xD3, 0xE2, 0xEC)


def textbox(slide, x, y, w, h, content, *, size=9.5, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.0,
            font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    para = tf.paragraphs[0]
    para.alignment = align
    para.line_spacing = line_spacing
    run = para.add_run()
    run.text = content
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    return box


def section(slide, y, title):
    """Draw a section header and rule; return nothing (body y is fixed per call)."""
    textbox(slide, M, y, W - 2 * M, Inches(0.2), title.upper(), size=8.6,
            bold=True, color=ACCENT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, y + Inches(0.19),
                                  W - 2 * M, Emu(9525))
    line.fill.solid()
    line.fill.fore_color.rgb = RULE
    line.line.fill.background()


def body(slide, y, h, items, *, size=9.4, space_before=3.5):
    """One text frame with all bullets: the text flows, nothing overlaps."""
    box = slide.shapes.add_textbox(M, y, W - 2 * M, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, item in enumerate(items):
        para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        para.space_before = Pt(space_before if index else 0)
        para.line_spacing = 1.0
        p_pr = para._p.get_or_add_pPr()
        p_pr.set("marL", "200000")
        p_pr.set("indent", "-200000")
        run = para.add_run()
        run.text = "•  " + item
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        run.font.color.rgb = INK


def table(slide, rows, x, y, w, h, *, col_widths, size=8.4):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, h)
    tbl = shape.table
    total = sum(col_widths)
    for idx, cw in enumerate(col_widths):
        tbl.columns[idx].width = Emu(int(w * cw / total))
    for r_i, row in enumerate(rows):
        for c_i, value in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            cell.text = ""
            tf = cell.text_frame
            tf.margin_left = Inches(0.06)
            tf.margin_right = Inches(0.05)
            tf.margin_top = Inches(0.01)
            tf.margin_bottom = Inches(0.01)
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if c_i == 0 else PP_ALIGN.RIGHT
            run = para.add_run()
            run.text = str(value)
            run.font.size = Pt(size)
            run.font.name = "Calibri"
            run.font.bold = (r_i == 0)
            run.font.color.rgb = WHITE if r_i == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r_i == 0 else (
                LIGHT if r_i % 2 == 0 else WHITE
            )
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def box_text(slide, x, y, w, h, content, *, fill=LIGHT, border=RULE, size=8.9,
             color=INK):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = border
    box.adjustments[0] = 0.07
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    para.line_spacing = 1.0
    run = para.add_run()
    run.text = content
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run.font.color.rgb = color


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Header band
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.2))
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    textbox(slide, M, Inches(0.24), W - 2 * M, Inches(0.45),
            "MCP + RAG for 1C — one-page summary", size=19, bold=True,
            color=WHITE)
    textbox(slide, M, Inches(0.76), W - 2 * M, Inches(0.35),
            "Local-first search over BSL business logic and configuration metadata  ·  "
            "Sergey Matyukhin",
            size=8.8, color=RGBColor(0xCF, 0xE4, 0xEF))

    # Problem
    section(slide, Inches(1.4), "Problem")
    body(slide, Inches(1.72), Inches(1.0), [
        "A 1C configuration dump is thousands of .bsl modules plus a multi-megabyte metadata report; "
        "grep finds text, not meaning.",
        "An exact handler name and an intent question need different search strategies, and an LLM agent "
        "needs a tool with measured quality.",
        "Hosted search stacks require deployment work and send the code off the machine.",
    ], size=9.5)

    # Approach
    section(slide, Inches(2.78), "Approach")
    body(slide, Inches(3.1), Inches(1.2), [
        "Local MCP server over stdio: 9 tools, including an explainable router and a diagnostic "
        "index_stats.",
        "intfloat/multilingual-e5-small (384-d, 512-token limit) + Chroma with HNSW cosine, CPU only.",
        "Custom global BM25 inverted index; Reciprocal Rank Fusion with k = 60; optional flashrank "
        "cross-encoder, off by default.",
        "No external APIs and no data leaving the machine; 40-60 ms per query.",
    ], size=9.5)

    # Corpus
    section(slide, Inches(4.48), "Corpus")
    body(slide, Inches(4.8), Inches(1.05), [
        "154,560 BSL code chunks (one procedure or function per chunk, ADR 0007) + 25,026 metadata chunks "
        "from the configuration XML = 179,586 chunks.",
        "Chunking gates: 99.6% code coverage; 0.3% of chunks over 512 tokens; token p50/p99/max = "
        "179 / 472 / 554.",
        "Metadata source change (text report → XML): BM25 Hit@3 93.0% → 99.7% (ADR 0008).",
    ], size=9.5)

    # Quality
    section(slide, Inches(6.0),
            "Quality — gold-set v2: 70 natural-language queries, file-level relevance, 4 runs, mean ± std")
    table(
        slide,
        [
            ["Configuration", "nDCG@10", "Recall@10", "Hit@1", "MRR@10"],
            ["vector", "0.339 ± 0.023", "0.523 ± 0.030", "0.150 ± 0.007", "0.268 ± 0.019"],
            ["pearson", "0.337 ± 0.030", "0.511 ± 0.048", "0.179 ± 0.021", "0.270 ± 0.024"],
            ["hybrid (RRF)", "0.301 ± 0.017", "0.482 ± 0.034", "0.129 ± 0.010", "0.240 ± 0.013"],
            ["hybrid + rerank", "0.117", "0.157", "0.086", "0.102"],
        ],
        M,
        Inches(6.34),
        W - 2 * M,
        Inches(1.12),
        col_widths=[2.5, 1.9, 1.9, 1.7, 1.7],
        size=8.4,
    )
    box_text(
        slide, M, Inches(7.58), W - 2 * M, Inches(1.48),
        "Paired test, vector vs hybrid: Δ nDCG@10 = +0.037, bootstrap 95% CI [-0.015; +0.086], "
        "p = 0.154 → NOT statistically significant. The default is unchanged; ADR 0009 is proposed.\n"
        "2×2 ablation (nDCG@10, vector/hybrid): old export · symbolic 0.176/0.209; old export · "
        "procedural not measurable (dump gone); current export · symbolic 0.165/0.205; current export · "
        "procedural 0.339/0.301. Chunking explains the leader flip; the export effect at symbolic "
        "chunking is −0.012/−0.004, inside run noise. Interaction cell missing, additivity assumed. "
        "The cross-encoder stays an option only (0.117).",
        size=8.5,
    )

    # Engineering
    section(slide, Inches(9.18), "Engineering")
    body(slide, Inches(9.5), Inches(1.4), [
        "72 unit tests; ruff clean, mypy over 22 modules, cyclomatic-complexity gate 143 blocks / "
        "avg 4.37 (A) in CI; Python 3.11/3.12; CI green at 7b50775.",
        "ADR log 0001-0009 and CHANGELOG; router 68/68 on the dev set (blind 75% before patches, 20/20 "
        "after — that set is dev now, so a new hold-out is needed).",
        "Limitations: HNSW rebuilt on start (~60-70 s cold start, workaround for chroma#7090); single "
        "writer; no incremental indexing; gold set partly synthetic.",
    ], size=9.5)

    # Links
    section(slide, Inches(10.68), "Links and contact")
    textbox(slide, M, Inches(11.0), W - 2 * M, Inches(0.3),
            "github.com/Smaturus/1c-mcp-rag-search   ·   smaturus@gmail.com   ·   "
            "MCP Python SDK · Chroma · multilingual-e5-small · BM25 · RRF · flashrank",
            size=8.4, bold=True, color=NAVY)

    textbox(slide, M, Inches(11.38), W - 2 * M, Inches(0.18),
            "All numbers come from the repository evaluation reports (code_proc / gold-set v2, 4 runs).",
            size=7.0, color=GRAY)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


def main() -> int:
    print("saved:", build())
    return 0


if __name__ == "__main__":
    sys.exit(main())
